"""Serial PDFium rendering/text and pypdf structural operations."""

import html
from contextlib import contextmanager
from pathlib import Path

import pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter

from ..core.formats import file_item, sorted_items


def reader(path):
    document = PdfReader(path)
    if document.is_encrypted:
        raise ValueError("Password-protected PDFs are unsupported in v1")
    return document


@contextmanager
def document(path):
    with Path(path).open("rb") as stream:
        checked = PdfReader(stream)
        if checked.is_encrypted:
            raise ValueError("Password-protected PDFs are unsupported in v1")
    pdf = pdfium.PdfDocument(path)
    try:
        pdf.init_forms()
        yield pdf
    finally:
        pdf.close()


def render_page(pdf, index, dpi):
    page = pdf[index]
    try:
        bitmap = page.render(scale=dpi / 72, fill_color=(255, 255, 255, 255), may_draw_forms=True)
        try:
            return bitmap.to_pil().copy()
        finally:
            bitmap.close()
    finally:
        page.close()


def extract_text(path, context):
    texts = []
    with document(path) as pdf:
        for index in range(len(pdf)):
            context.token.check()
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    texts.append(textpage.get_text_range().replace("\r\n", "\n").strip())
                finally:
                    textpage.close()
            finally:
                page.close()
    if not any(texts):
        raise ValueError("No selectable text found. Use Recognize text (OCR) for scanned PDFs")
    return texts


def page_name(item, index, count, target):
    return f"{item.basename} - {index + 1:0{len(str(count))}d}.{target}"


def render(item, target, context, source_path=None):
    with document(source_path or item.path) as pdf:
        if len(pdf) == 1:
            return [context.produce(item, target, lambda path: save_page(pdf, 0, path, target, context))]
        stage = context.publisher.stage(item.path.parent, directory=True)
        for index in range(len(pdf)):
            context.token.check()
            save_page(pdf, index, stage / page_name(item, index, len(pdf), target), target, context)
            context.progress((index + 1) / len(pdf), f"Page {index + 1}")
        return [context.publisher.publish(stage, item.path.parent, item.basename + " (pages)", "", context.token)]


def save_page(pdf, index, path, target, context):
    with render_page(pdf, index, context.settings.pdf_dpi) as image:
        options = {"quality": context.settings.jpeg_quality} if target == "jpg" else {}
        image.convert("RGB").save(path, "JPEG" if target == "jpg" else "PNG", **options)


def split(item, context):
    pdf = reader(item.path)
    stage = context.publisher.stage(item.path.parent, directory=True)
    for index, page in enumerate(pdf.pages):
        context.token.check()
        with PdfWriter() as writer:
            writer.add_page(page)
            writer.write(stage / page_name(item, index, len(pdf.pages), "pdf"))
    return [context.publisher.publish(stage, item.path.parent, item.basename + " (pages)", "", context.token)]


def structural(op, items, context):
    ordered = sorted_items(items)
    suffix = {"pdf.merge": "merged", "pdf.rotate": "rotated", "pdf.optimize": "optimized"}[op.id]
    def write(path):
        with PdfWriter() as writer:
            for item in ordered:
                context.token.check()
                pdf = reader(item.path)
                for page in pdf.pages:
                    if op.id == "pdf.rotate":
                        page.rotate(90)
                    writer.add_page(page)
            if op.id == "pdf.optimize":
                optimize(writer, context)
            writer.write(path)
    return [context.produce(ordered[0], "pdf", write, suffix)]


def optimize(writer, context):
    for page in writer.pages:
        context.token.check()
        page.compress_content_streams()
    writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
    context.result.notes.append("Structural optimization does not downsample images or guarantee a smaller PDF")


def to_docx(texts, path):
    from docx import Document
    output = Document()
    for index, text in enumerate(texts):
        if index:
            output.add_page_break()
        for paragraph in text.split("\n"):
            output.add_paragraph(paragraph)
    output.save(path)


def to_pptx(item, path, context):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    with document(item.path) as pdf:
        first = pdf[0]
        try:
            ratio = first.get_width() / first.get_height()
        finally:
            first.close()
        presentation.slide_width = Inches(10)
        presentation.slide_height = int(presentation.slide_width / ratio)
        for index in range(len(pdf)):
            context.token.check()
            image_path = context.workspace / f"slide-{index}.png"
            save_page(pdf, index, image_path, "png", context)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            fit_slide(slide, image_path, presentation)
    presentation.save(path)


def fit_slide(slide, image_path, presentation):
    from PIL import Image
    with Image.open(image_path) as image:
        ratio = image.width / image.height
    width = min(presentation.slide_width, int(presentation.slide_height * ratio))
    height = int(width / ratio)
    slide.shapes.add_picture(str(image_path), (presentation.slide_width - width) // 2, (presentation.slide_height - height) // 2, width=width, height=height)


def text_document(item, target, context):
    texts = extract_text(item.path, context)
    if target == "docx":
        context.result.notes.append("Text-only DOCX; PDF layout is not reconstructed")
        return [context.produce(item, target, lambda path: to_docx(texts, path))]
    if target in {"txt", "md"}:
        separator = "\n\n---\n\n" if target == "md" else "\n\n\f\n\n"
        return [context.produce(item, target, lambda path: path.write_text(separator.join(texts), encoding="utf-8"))]
    from .pandoc import convert_file
    source = context.workspace / "extracted.html"
    source.write_text("<html><body>" + "<hr>".join("<pre>" + html.escape(text) + "</pre>" for text in texts) + "</body></html>", encoding="utf-8")
    return [context.produce(item, target, lambda path: convert_file(source, "html", path, target, context))]


def run(op, items, context):
    item = items[0]
    if op.id in {"pdf.merge", "pdf.rotate", "pdf.optimize"}:
        return structural(op, items, context)
    if op.id == "pdf.split":
        return split(item, context)
    target = op.target_format
    if target in {"png", "jpg"}:
        return render(item, target, context)
    if target == "pptx":
        return [context.produce(item, target, lambda path: to_pptx(item, path, context))]
    if target in {"docx", "txt", "md", "html", "epub"}:
        return text_document(item, target, context)
    raise ValueError(f"Unsupported PDF target: {target}")
