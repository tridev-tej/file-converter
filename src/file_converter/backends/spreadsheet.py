"""Values-oriented workbook readers and XLSX/delimited writers."""

import csv
import datetime
import io
import math
from dataclasses import dataclass
from pathlib import Path

from ..core.formats import safe_name


@dataclass
class Sheet:
    name: str
    rows: list[list]


def trim_rows(rows):
    result = [list(row) for row in rows]
    while result and all(value is None or value == "" for value in result[-1]):
        result.pop()
    return result


def read_openpyxl(path):
    from openpyxl import load_workbook
    values = load_workbook(path, data_only=True, read_only=True)
    formulas = load_workbook(path, data_only=False, read_only=True)
    result = []
    try:
        for sheet in formulas:
            cached = values[sheet.title]
            rows = []
            for row in sheet:
                rows.append([cached_value(cell, cached) for cell in row])
            result.append(Sheet(sheet.title, trim_rows(rows)))
    finally:
        values.close()
        formulas.close()
    return result


def cached_value(cell, cached):
    if cell.data_type == "f":
        value = cached.cell(cell.row, cell.column).value
        if value is None:
            raise ValueError(f"Formula has no cached value: {cell.parent.title}!{cell.coordinate}. Recalculate and save the workbook first")
        return value
    value = cell.value
    if isinstance(value, (int, float)) and len(str(abs(value)).replace(".", "")) > 15:
        return str(value)
    if isinstance(value, (int, float)) and cell.number_format and set(cell.number_format) == {"0"}:
        return str(int(value)).zfill(len(cell.number_format))
    return value


def read_xls(path):
    import xlrd
    book = xlrd.open_workbook(path)
    result = []
    try:
        for sheet in book.sheets():
            rows = [[xls_value(sheet.cell(row, col), book.datemode) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
            result.append(Sheet(sheet.name, trim_rows(rows)))
    finally:
        book.release_resources()
    return result


def xls_value(cell, datemode):
    import xlrd
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return None
    if cell.ctype == xlrd.XL_CELL_ERROR:
        raise ValueError("XLS cell contains an error")
    return cell.value


def ods_value(cell):
    from odf import teletype
    kind = cell.getAttribute("valuetype")
    if cell.getAttribute("formula") and not kind:
        raise ValueError("ODS formula has no cached value")
    if kind in {"float", "currency", "percentage"}:
        value = cell.getAttribute("value")
        if value is None:
            raise ValueError("ODS numeric cell has no cached value")
        return float(value)
    if kind == "boolean":
        return cell.getAttribute("booleanvalue") == "true"
    if kind == "date":
        return datetime.datetime.fromisoformat(cell.getAttribute("datevalue"))
    return teletype.extractText(cell) or None


def ods_row(row):
    from odf.table import TableCell, CoveredTableCell
    values = []
    for cell in row.childNodes:
        if getattr(cell, "qname", None) not in {TableCell().qname, CoveredTableCell().qname}:
            continue
        repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        if repeat > 16384:
            raise ValueError("ODS repeated column count exceeds supported limits")
        values.extend([ods_value(cell)] * repeat)
    while values and values[-1] is None:
        values.pop()
    return values


def read_ods(path):
    from odf.opendocument import load
    from odf.table import Table, TableRow
    document = load(str(path))
    sheets = []
    for table in document.spreadsheet.getElementsByType(Table):
        rows = []
        pending_empty = 0
        for row in table.getElementsByType(TableRow):
            values = ods_row(row)
            repeat = int(row.getAttribute("numberrowsrepeated") or 1)
            if repeat > 1048576 or len(rows) + repeat > 1048576:
                raise ValueError("ODS repeated row count exceeds supported limits")
            if values:
                rows.extend([[] for _ in range(pending_empty)])
                pending_empty = 0
                rows.extend([values.copy() for _ in range(repeat)])
            else:
                pending_empty += repeat
        sheets.append(Sheet(table.getAttribute("name") or "Sheet", trim_rows(rows)))
    return sheets


def detect_delimiter(text):
    sample = text[:65536]
    candidates = []
    for delimiter in (",", ";", "\t", "|"):
        if delimiter_matches(sample, delimiter):
            candidates.append(delimiter)
    if len(candidates) == 1:
        return candidates[0]
    if not candidates and not any(delimiter in sample for delimiter in (",", ";", "\t", "|")):
        return ","
    raise ValueError("CSV delimiter is ambiguous; use a consistently delimited UTF-8 file")


def delimiter_matches(sample, delimiter):
    rows = list(csv.reader(io.StringIO(sample), delimiter=delimiter))[:20]
    widths = [len(row) for row in rows if row]
    return bool(widths) and min(widths) > 1 and len(set(widths)) == 1


def read_delimited(path, fmt):
    text = Path(path).read_text(encoding="utf-8-sig")
    delimiter = "\t" if fmt == "tsv" else detect_delimiter(text)
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True))
    return [Sheet("Sheet1", trim_rows(rows))]


def read_workbook(path, fmt):
    readers = {"xlsx": read_openpyxl, "xlsm": read_openpyxl, "xls": read_xls, "ods": read_ods}
    if fmt in {"csv", "tsv"}:
        return read_delimited(path, fmt)
    if fmt not in readers:
        raise ValueError(f"Unsupported spreadsheet input: {fmt}")
    return readers[fmt](path)


def write_xlsx(sheets, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for sheet in sheets:
        target = book.create_sheet(sheet.name[:31])
        for row in sheet.rows:
            target.append(row)
            preserve_string_cells(target[target.max_row])
        for cell in target[1]:
            cell.font = Font(bold=True)
        widths = column_widths(sheet.rows)
        for col, width in enumerate(widths, 1):
            target.column_dimensions[get_column_letter(col)].width = width
    book.save(path)
    book.close()


def preserve_string_cells(cells):
    for cell in cells:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def column_widths(rows):
    widths = []
    for row in rows:
        for index, value in enumerate(row):
            if index >= len(widths):
                widths.append(10)
            widths[index] = min(60, max(widths[index], len(str(value or "")) + 2))
    return widths


def write_delimited(sheet, path, target):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter="\t" if target == "tsv" else ",", lineterminator="\r\n")
        for row in sheet.rows:
            writer.writerow(row)


def run(op, items, context):
    item = items[0]
    target = op.target_format
    if target in {"ods", "pdf"}:
        from .office import convert_to_output
        return convert_to_output(item, target, context)
    sheets = read_workbook(item.path, item.format)
    if not sheets:
        raise ValueError("Workbook contains no sheets")
    context.result.notes.append("Values only: macros, charts, styling and formula calculations are not preserved")
    if target == "xlsx":
        return [context.produce(item, "xlsx", lambda path: write_xlsx(sheets, path))]
    if target not in {"csv", "tsv"}:
        raise ValueError(f"Unsupported spreadsheet target: {target}")
    populated = [sheet for sheet in sheets if sheet.rows]
    if not populated:
        populated = sheets[:1]
        context.result.notes.append("Workbook is empty; emitted the first sheet")
    return export_sheets(item, populated, target, context)


def export_sheets(item, sheets, target, context):
    from dataclasses import replace
    outputs = []
    for sheet in sheets:
        context.token.check()
        base = item.basename if len(sheets) == 1 else item.basename + " - " + safe_name(sheet.name)
        named = replace(item, basename=base)
        output = context.produce(named, target, lambda path: write_delimited(sheet, path, target))
        context.result.outputs.append(output)
    return outputs
