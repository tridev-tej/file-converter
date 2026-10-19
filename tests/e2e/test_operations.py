import json
import subprocess
import sys
import time
from pathlib import Path

import pytest
from conftest import REPORT

IDS = Path(__file__).with_name("expected_ids.txt").read_text().splitlines()


def inputs_for(op, root):
    prefix = op.id.split(".")[0]
    names = {"image": ["one.png"], "gif": ["animation.gif"], "audio": ["one.wav"], "video": ["clip.mp4"], "pdf": ["one.pdf"], "doc": ["notes.md"], "slides": ["slides.pptx"], "sheet": ["book.xlsx"], "subs": ["captions.srt"], "archive": ["sample.zip"], "pack": ["folder"], "tools": ["one.png"], "action": ["folder"]}
    same = {"image.to.png": ["two.jpg"], "audio.to.wav": ["one.mp3"], "video.to.mp4": ["animation.gif"], "doc.to.md": ["notes.txt"], "sheet.to.xlsx": ["data.csv"], "subs.to.srt": ["captions.vtt"], "archive.to.zip": ["raw.tar"], "tools.ocr": ["ocr.png"], "slides.to.pptx": ["slides.odp"]}
    multiple = {"image.merge.pdf": ["one.png", "two.jpg"], "pdf.merge": ["two.pdf", "one.pdf"], "audio.concat": ["two.wav", "one.wav"]}
    chosen = multiple.get(op.id, same.get(op.id, names[prefix]))
    return [root / name for name in chosen]


def ensure_extra(root):
    import tarfile
    from pptx import Presentation
    if not (root / "raw.tar").exists():
        with tarfile.open(root / "raw.tar", "w") as archive:
            archive.add(root / "folder", arcname="folder")
        slides = Presentation()
        slides.slides.add_slide(slides.slide_layouts[6])
        slides.save(root / "slides.pptx")
        from odf.opendocument import OpenDocumentPresentation
        from odf.draw import Page, Frame, TextBox
        from odf.text import P
        odp = OpenDocumentPresentation()
        page = Page(name="Slide1", masterpagename="Default")
        frame = Frame(width="10cm", height="3cm", x="1cm", y="1cm")
        box = TextBox()
        box.addElement(P(text="FileConverter slide"))
        frame.addElement(box)
        page.addElement(frame)
        odp.presentation.addElement(page)
        odp.save(str(root / "slides.odp"))


def inspect_output(path, op):
    from PIL import Image
    from pypdf import PdfReader
    from file_converter.core.formats import file_item
    from file_converter.backends.spreadsheet import read_workbook
    from file_converter.backends.subtitle import parse
    fmt = file_item(path).format
    if path.is_dir():
        assert list(path.rglob("*"))
        return
    assert path.stat().st_size > 0
    if fmt in {"jpg", "png", "webp", "avif", "tiff", "ico", "gif", "heic"}:
        with Image.open(path) as image:
            image.load()
            assert min(image.size) > 0
    elif fmt == "pdf":
        assert len(PdfReader(path).pages) > 0
    elif fmt in {"xlsx", "csv", "tsv"}:
        assert read_workbook(path, fmt)
    elif fmt in {"srt", "vtt"}:
        assert parse(path)
    elif fmt in {"mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "mp4", "mov", "mkv", "webm"}:
        import imageio_ffmpeg
        command = [imageio_ffmpeg.get_ffmpeg_exe(), "-v", "error", "-i", str(path), "-f", "null", "-"]
        subprocess.run(command, check=True, capture_output=True)
    elif fmt == "docx":
        from docx import Document
        assert "FileConverter" in "\n".join(p.text for p in Document(path).paragraphs)
    else:
        inspect_document(path, fmt)


def inspect_document(path, fmt):
    if fmt in {"zip", "epub", "pptx", "odt"}:
        import zipfile
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
    elif fmt in {"txt", "md", "html", "rtf"}:
        assert path.read_text(encoding="utf-8-sig").strip()


@pytest.mark.parametrize("operation_id", IDS)
def test_operation(operation_id, samples):
    from file_converter.core.catalog import build_catalog
    from file_converter.core.formats import file_item
    from file_converter.core.runner import execute
    from file_converter.core.operations import OperationRequest
    from file_converter.core.settings import Settings
    from file_converter.core.tools import discover_tools
    ensure_extra(samples)
    op = next(op for op in build_catalog() if op.id == operation_id)
    inputs = inputs_for(op, samples)
    inputs = prepare_video_input(op, samples, inputs)
    items = tuple(file_item(path) for path in inputs)
    available = op.availability(items, discover_tools())
    record = {"id": op.id, "inputs": list(map(str, inputs)), "route": available.route, "status": "failed", "outputs": [], "assertions": ["public execution result", "readable output content"], "duration": 0}
    REPORT["source"].append(record)
    if not available.enabled:
        record.update(status="skipped", error="; ".join(available.reasons))
        pytest.skip(record["error"])
    if op.id in {"action.path", "action.terminal", "action.reveal"}:
        record.update(status="manual_required", error="Desktop effect or real Vision acceptance requires a desktop session")
        pytest.skip(record["error"])
    params = {"start": "0.1", "end": "0.8"} if op.id.endswith("trim") else {}
    if op.id in {"action.copy", "action.move"}:
        import shutil
        source = samples / (op.id + " source")
        shutil.copytree(samples / "folder", source)
        inputs = [source]
        params = {"destination": str(samples / "destination")}
    started = time.monotonic()
    result = execute(OperationRequest(op.id, tuple(inputs), params, Settings()))
    record.update(status="failed", execution_state=result.state, outputs=list(map(str, result.outputs)), error=result.errors, duration=time.monotonic()-started)
    assert result.state == "succeeded", result.errors
    assert result.outputs
    assert_ocr(op, result)
    for output in result.outputs:
        inspect_output(output, op)
    record["status"] = "passed"


def test_cli_contract(samples):
    import os
    env = dict(os.environ, PYTHONPATH=str(Path("src").absolute()))
    command = [sys.executable, "-m", "file_converter"]
    listed = subprocess.run([*command, "list", str(samples / "one.png"), "--json"], capture_output=True, text=True, env=env)
    assert listed.returncode == 0, listed.stderr
    assert "image.to.jpg" in [row["id"] for row in json.loads(listed.stdout)]
    converted = subprocess.run([*command, "run", "image.to.jpg", str(samples / "one.png"), "--json"], capture_output=True, text=True, env=env)
    assert converted.returncode == 0, converted.stderr
    assert json.loads(converted.stdout)["state"] == "succeeded"


def test_cli_move(tmp_path):
    import os
    source = tmp_path / "move me.txt"
    source.write_text("moved content")
    destination = tmp_path / "destination"
    destination.mkdir()
    env = dict(os.environ, PYTHONPATH=str(Path("src").absolute()))
    command = [sys.executable, "-m", "file_converter", "run", "action.move", str(source), "--param", "destination=" + str(destination), "--json"]
    result = subprocess.run(command, capture_output=True, text=True, env=env)
    assert result.returncode == 0, result.stdout + result.stderr
    data = json.loads(result.stdout)
    assert data["state"] == "succeeded"
    assert Path(data["outputs"][0]).read_text() == "moved content"
    assert not source.exists()


def test_adversarial_contract(samples, tmp_path):
    import zipfile
    from file_converter.core.operations import OperationRequest
    from file_converter.core.runner import execute, validate_request
    from file_converter.core.settings import Settings, load_settings
    from file_converter.backends.ffmpeg import parse_time
    settings = Settings()
    for value in ["nan", "inf", "-1", "1:99", "1:2:99"]:
        with pytest.raises(ValueError):
            parse_time(value)
    with pytest.raises(ValueError):
        validate_request(OperationRequest("image.to.jpg", (samples / "one.png", samples / "notes.txt"), {}, settings))
    bad = tmp_path / "bad.zip"
    with zipfile.ZipFile(bad, "w") as archive:
        archive.writestr("../escaped.txt", "bad")
    result = execute(OperationRequest("archive.extract", (bad,), {}, settings))
    assert result.state == "failed"
    assert not (tmp_path.parent / "escaped.txt").exists()
    result = execute(OperationRequest("pdf.to.txt", (samples / "scan.pdf",), {}, settings))
    assert result.state == "failed"
    result = execute(OperationRequest("video.to.mp3", (samples / "silent.mp4",), {}, settings))
    assert result.state == "failed"
    config = tmp_path / "settings.json"
    config.write_text("{")
    with pytest.raises(ValueError):
        load_settings(config)


def prepare_video_input(op, samples, inputs):
    if op.id != "video.to.mp4":
        return inputs
    import shutil
    output = samples / "other.mov"
    shutil.copyfile(samples / "clip.mp4", output)
    return [output]


def assert_ocr(op, result):
    if op.id == "tools.ocr":
        assert "FILE CONVERTER TEST 123" in result.outputs[0].read_text().upper()
